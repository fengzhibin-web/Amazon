import json
import math
import os
import re
import subprocess
import sys
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    DND_AVAILABLE = True
except Exception:
    DND_AVAILABLE = False
    DND_FILES = None
    TkinterDnD = None

APP_TITLE = "Excel 商品数据提取工具"
WINDOW_SIZE = "1240x800"
INVISIBLE_CHAR_PATTERN = re.compile(r"[\u200b\u200c\u200d\ufeff]")

REQUIRED_GEMINI_FIELDS = [
    "id", "品牌识别", "AI_通用标题_EN", "AI_通用标题_CN",
    "AI_卖点_EN_1", "AI_卖点_CN_1", "AI_卖点_EN_2", "AI_卖点_CN_2",
    "AI_卖点_EN_3", "AI_卖点_CN_3", "AI_卖点_EN_4", "AI_卖点_CN_4",
    "AI_卖点_EN_5", "AI_卖点_CN_5", "AI_规格描述_EN", "AI_规格描述_CN",
]

BASE_VERTICAL_FIELDS = [
    "id", "品牌识别", "AI_通用标题_EN", "AI_通用标题_CN",
    "AI_卖点_EN_1", "AI_卖点_EN_2", "AI_卖点_EN_3", "AI_卖点_EN_4", "AI_卖点_EN_5",
    "AI_卖点_CN_1", "AI_卖点_CN_2", "AI_卖点_CN_3", "AI_卖点_CN_4", "AI_卖点_CN_5",
    "AI_规格描述_EN", "AI_规格描述_CN",
]

OLD_SPEC_EN_FIELDS = [f"AI_规格描述_EN_{i}" for i in range(1, 6)]
OLD_SPEC_CN_FIELDS = [f"AI_规格描述_CN_{i}" for i in range(1, 6)]

SELLING_POINT_FIELDS = {
    "AI_卖点_EN_1", "AI_卖点_EN_2", "AI_卖点_EN_3", "AI_卖点_EN_4", "AI_卖点_EN_5",
    "AI_卖点_CN_1", "AI_卖点_CN_2", "AI_卖点_CN_3", "AI_卖点_CN_4", "AI_卖点_CN_5",
}


class ExcelExtractorApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry(WINDOW_SIZE)

        self.input_path_var = tk.StringVar()
        self.batch_size_var = tk.StringVar(value="10")

        self.dedup_output_dir_var = tk.StringVar()
        self.dedup_filename_var = tk.StringVar(value="extracted_result.xlsx")
        self.batch_output_dir_var = tk.StringVar()

        self.vertical_output_dir_var = tk.StringVar()
        self.vertical_filename_var = tk.StringVar(value=self._default_vertical_filename())

        self.batch_source_txt_var = tk.StringVar()
        self.include_original_title_var = tk.BooleanVar(value=False)
        self.include_original_desc_var = tk.BooleanVar(value=False)

        self.tasks = []
        self.task_counter = 0
        self._log_link_counter = 0
        self._clean_stats = {"cite": 0, "tail_blank": 0, "wrap_quote": 0}
        self._batch_txt_entry = None
        self._batch_txt_btn = None

        self._build_ui()

    def _build_ui(self):
        top_frame = tk.Frame(self.root)
        top_frame.pack(fill="both", expand=True, padx=12, pady=8)

        tk.Label(
            top_frame,
            text="流程1：原始Excel -> 去重导出 + Gemini批量TXT | 流程2：Gemini结果(TSV/文件) -> 纵向Excel",
            font=("Microsoft YaHei", 10, "bold"), anchor="w"
        ).pack(fill="x", pady=(0, 6))

        panes = tk.PanedWindow(top_frame, orient="horizontal", sashrelief="raised")
        panes.pack(fill="both", expand=True)

        left = tk.Frame(panes)
        right = tk.Frame(panes)
        panes.add(left, minsize=780)
        panes.add(right, minsize=420)

        self._build_left_area(left)
        self._build_task_area(right)

    def _build_left_area(self, parent):
        notebook = ttk.Notebook(parent)
        notebook.pack(fill="both", expand=True)

        tab_step1 = tk.Frame(notebook)
        tab_step2 = tk.Frame(notebook)
        notebook.add(tab_step1, text="流程1：去重+切分")
        notebook.add(tab_step2, text="流程2：Gemini转纵向")

        self._build_tab_step1(tab_step1)
        self._build_tab_step2(tab_step2)

        tk.Label(parent, text="运行信息（可点击“打开目录”快速跳转）", anchor="w").pack(fill="x", pady=(8, 4))
        self.log_text = tk.Text(parent, height=12, wrap="word", state="disabled")
        self.log_text.pack(fill="both", expand=True)

    def _build_tab_step1(self, parent):
        tutorial = (
            "流程1教程：\n1) 选择原始Excel。\n2) 配置批量大小（一个txt包含多少ID）。\n"
            "3) 配置输出目录与文件名。\n4) 点击运行，自动完成去重+切分。"
        )
        tk.Label(parent, text=tutorial, justify="left", fg="#334155").pack(fill="x", pady=(0, 8))

        self._row_file(parent, "原始 Excel:", self.input_path_var, self.choose_input_file, {".xlsx", ".xls"})

        row = tk.Frame(parent)
        row.pack(fill="x", pady=4)
        tk.Label(row, text="批量大小:", width=22, anchor="w").pack(side="left")
        tk.Entry(row, textvariable=self.batch_size_var, width=10).pack(side="left")
        tk.Label(row, text="（一个文件包含多少ID）", fg="#475569").pack(side="left", padx=(8, 0))

        self._row_dir(parent, "去重输出目录:", self.dedup_output_dir_var, self.choose_dedup_output_dir)
        self._row_file_name(parent, "去重输出文件名:", self.dedup_filename_var)
        self._row_dir(parent, "批量TXT输出目录:", self.batch_output_dir_var, self.choose_batch_output_dir)

        self._add_drop_zone(parent, "拖拽上传区（流程1 Excel）", self.input_path_var, {".xlsx", ".xls"})

        tk.Button(parent, text="运行流程1（去重+切分）", width=24, bg="#0ea5e9", fg="white", command=self.run_step1_pipeline).pack(anchor="w", pady=10)

    def _build_tab_step2(self, parent):
        tutorial = (
            "流程2教程：\n直接粘贴网页 Gemini TSV 文本后转换。"
        )
        tk.Label(parent, text=tutorial, justify="left", fg="#334155").pack(fill="x", pady=(0, 8))

        cb_row = tk.Frame(parent)
        cb_row.pack(fill="x", pady=4)
        tk.Checkbutton(cb_row, text="输出原始标题", variable=self.include_original_title_var, command=self._toggle_batch_source_widgets).pack(side="left")
        tk.Checkbutton(cb_row, text="输出原始描述", variable=self.include_original_desc_var, command=self._toggle_batch_source_widgets).pack(side="left", padx=(12, 0))

        row = tk.Frame(parent)
        row.pack(fill="x", pady=4)
        tk.Label(row, text="Batch源文件TXT:", width=22, anchor="w").pack(side="left")
        self._batch_txt_entry = tk.Entry(row, textvariable=self.batch_source_txt_var, state="disabled")
        self._batch_txt_entry.pack(side="left", fill="x", expand=True)
        self._batch_txt_btn = tk.Button(row, text="选择", width=10, command=self.choose_batch_source_txt, state="disabled")
        self._batch_txt_btn.pack(side="left", padx=(8, 0))
        if DND_AVAILABLE:
            self._enable_drop(self._batch_txt_entry, self.batch_source_txt_var, {".txt"})

        self._row_dir(parent, "纵向Excel输出目录:", self.vertical_output_dir_var, self.choose_vertical_output_dir)
        self._row_file_name(parent, "纵向输出文件名:", self.vertical_filename_var)

        tk.Label(parent, text="粘贴 Gemini TSV 文本（支持 ```tsv 包裹）:", anchor="w").pack(fill="x", pady=(8, 4))
        self.tsv_paste_text = tk.Text(parent, height=14, wrap="word")
        self.tsv_paste_text.pack(fill="both", expand=True)

        tk.Button(parent, text="从粘贴文本转换", width=18, bg="#9333ea", fg="white", command=self.run_convert_gemini_vertical_from_paste).pack(side="left", pady=8)

    def _build_task_area(self, parent):
        tk.Label(parent, text="任务列表", font=("Microsoft YaHei", 10, "bold")).pack(anchor="w", pady=(0, 6))
        columns = ("task_name", "source", "status", "count", "output", "created", "finished", "remark")
        tree = ttk.Treeview(parent, columns=columns, show="headings", height=32)
        headers = {"task_name": "任务名", "source": "输入来源", "status": "状态", "count": "产品数量", "output": "输出文件", "created": "创建时间", "finished": "完成时间", "remark": "备注"}
        widths = {"task_name": 120, "source": 100, "status": 70, "count": 70, "output": 140, "created": 120, "finished": 120, "remark": 280}
        for col in columns:
            tree.heading(col, text=headers[col])
            tree.column(col, width=widths[col], anchor="w")
        tree.pack(fill="both", expand=True)
        self.task_tree = tree

    def _row_file(self, parent, label, var, cmd, drop_types=None):
        row = tk.Frame(parent)
        row.pack(fill="x", pady=4)
        tk.Label(row, text=label, width=22, anchor="w").pack(side="left")
        entry = tk.Entry(row, textvariable=var)
        entry.pack(side="left", fill="x", expand=True)
        tk.Button(row, text="选择", width=10, command=cmd).pack(side="left", padx=(8, 0))
        if DND_AVAILABLE and drop_types:
            self._enable_drop(entry, var, drop_types)

    def _row_dir(self, parent, label, var, cmd):
        self._row_file(parent, label, var, cmd)

    def _row_file_name(self, parent, label, var):
        row = tk.Frame(parent)
        row.pack(fill="x", pady=4)
        tk.Label(row, text=label, width=22, anchor="w").pack(side="left")
        tk.Entry(row, textvariable=var).pack(side="left", fill="x", expand=True)

    def _add_drop_zone(self, parent, title, var, allowed_exts):
        zone = tk.Label(parent, text=title, relief="groove", bd=1, height=3, anchor="center", bg="#F8FAFC", fg="#334155")
        zone.pack(fill="x", pady=(4, 8))
        if DND_AVAILABLE:
            try:
                zone.drop_target_register(DND_FILES)
                zone.dnd_bind("<<Drop>>", lambda e: self._handle_drop(e, var, allowed_exts))
                zone.configure(text=title + "\n(可拖拽文件到此区域)")
            except Exception:
                zone.configure(text=title + "\n(拖拽初始化失败，请使用选择按钮)")
        else:
            zone.configure(text=title + "\n(当前环境未启用拖拽，请使用选择按钮)")

    def _enable_drop(self, widget, var_obj, allowed_exts):
        try:
            widget.drop_target_register(DND_FILES)
            widget.dnd_bind("<<Drop>>", lambda e: self._handle_drop(e, var_obj, allowed_exts))
        except Exception:
            self._append_log("提示：拖拽功能初始化失败，将继续使用选择按钮。")

    def _handle_drop(self, event, var_obj, allowed_exts):
        raw = event.data.strip()
        path = raw.strip("{}").split(" ")[0].strip("{}")
        ext = os.path.splitext(path)[1].lower()
        if ext not in allowed_exts:
            messagebox.showwarning("格式不支持", f"仅支持: {', '.join(sorted(allowed_exts))}")
            return
        var_obj.set(path)

    def _toggle_batch_source_widgets(self):
        enabled = self.include_original_title_var.get() or self.include_original_desc_var.get()
        state = "normal" if enabled else "disabled"
        if self._batch_txt_entry:
            self._batch_txt_entry.configure(state=state)
        if self._batch_txt_btn:
            self._batch_txt_btn.configure(state=state)

    def _append_log(self, msg: str, path: str | None = None):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"[{datetime.now().strftime('%H:%M:%S')}] {msg}")
        if path:
            self._log_link_counter += 1
            tag = f"link_{self._log_link_counter}"
            self.log_text.insert("end", "  [打开目录]", tag)
            self.log_text.tag_config(tag, foreground="#2563eb", underline=True)
            self.log_text.tag_bind(tag, "<Button-1>", lambda e, p=path: self._open_folder(p))
        self.log_text.insert("end", "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _open_folder(self, path: str):
        target = path if os.path.isdir(path) else os.path.dirname(path)
        try:
            if os.name == "nt":
                os.startfile(target)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", target])
            else:
                subprocess.Popen(["xdg-open", target])
        except Exception as exc:
            messagebox.showerror("错误", f"无法打开目录: {exc}")

    @staticmethod
    def _normalize_col_name(name: str) -> str:
        return INVISIBLE_CHAR_PATTERN.sub("", str(name)).strip()

    @staticmethod
    def _clean_cell(value):
        text = "" if value is None else str(value)
        text = text.strip().rstrip()
        if len(text) >= 2 and text[0] == '"' and text[-1] == '"':
            text = text[1:-1]
        return text.strip()

    @staticmethod
    def clean_leading_numbering(text):
        s = "" if text is None else str(text).strip()
        if not s:
            return ""
        return re.sub(r"^\s*(?:\(?\d+\)?|（\d+）)\s*(?:[\.、\):：\-]|-)?\s*", "", s).strip()

    def clean_export_text(self, text):
        if text is None:
            return ""
        original = str(text)
        val = original.replace("\r\n", "\n").replace("\r", "\n")

        cite_pattern = re.compile(r"\[cite:\s*[\d,\s]+\]", re.IGNORECASE)
        cite_count = len(cite_pattern.findall(val))
        if cite_count:
            self._clean_stats["cite"] += cite_count
        val = cite_pattern.sub("", val)

        val = re.sub(r"\[citation needed\]", "", val, flags=re.IGNORECASE)
        val = re.sub(r"【citation needed】", "", val, flags=re.IGNORECASE)
        val = re.sub(r"\[cite\s*start\]", "", val, flags=re.IGNORECASE)
        val = re.sub(r"\[cite\s*end\]", "", val, flags=re.IGNORECASE)

        val = re.sub(r"\s*\[\d+\](?=\s*(?:$|\n|。|\.|,|，|;|；))", "", val)

        val = val.strip()
        if len(val) >= 2 and val.startswith('"') and val.endswith('"'):
            self._clean_stats["wrap_quote"] += 1
            val = val[1:-1].strip()

        lines = val.split("\n")
        cleaned_lines = []
        for line in lines:
            line = line.strip()
            if not line:
                continue
            line = cite_pattern.sub("", line).strip()
            line = re.sub(r"\[cite\s*start\]|\[cite\s*end\]", "", line, flags=re.IGNORECASE).strip()
            line = re.sub(r"\s*\[\d+\]\s*$", "", line).strip()
            if line.endswith('"') and not line.startswith('"'):
                line = line[:-1].rstrip()
            if line.startswith('"') and not line.endswith('"'):
                line = line[1:].lstrip()
            if line:
                cleaned_lines.append(line)

        out = "\n".join(cleaned_lines).rstrip()
        if original.endswith("\n") or original.endswith("\r"):
            self._clean_stats["tail_blank"] += 1
        return out

    @staticmethod
    def clean_gemini_code_block(content: str) -> str:
        text = content.replace("\r\n", "\n").replace("\r", "\n")
        fence = re.search(r"```(?:tsv)?\s*\n(.*?)\n```", text, re.S | re.I)
        if fence:
            return fence.group(1).strip()
        return "\n".join([ln for ln in text.split("\n") if ln.strip()]).strip()

    @staticmethod
    def _default_vertical_filename() -> str:
        return f"gemini_result_vertical_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    def _ensure_xlsx_name(self, filename: str) -> str:
        name = filename.strip() or self._default_vertical_filename()
        return name if name.lower().endswith(".xlsx") else name + ".xlsx"

    def _create_task(self, task_name, source):
        self.task_counter += 1
        task = {
            "id": f"task_{self.task_counter:04d}", "task_name": task_name, "source": source,
            "status": "处理中", "count": "", "output": "", "created": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "finished": "", "remark": "",
        }
        self.tasks.append(task)
        self.task_tree.insert("", "end", iid=task["id"], values=(task["task_name"], task["source"], task["status"], task["count"], task["output"], task["created"], task["finished"], task["remark"]))
        return task

    def _update_task(self, task, status=None, count=None, output=None, remark=None):
        if status is not None:
            task["status"] = status
        if count is not None:
            task["count"] = count
        if output is not None:
            task["output"] = output
        if remark is not None:
            task["remark"] = remark
        if status in {"已完成", "失败"}:
            task["finished"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.task_tree.item(task["id"], values=(task["task_name"], task["source"], task["status"], task["count"], task["output"], task["created"], task["finished"], task["remark"]))

    def choose_input_file(self):
        file_path = filedialog.askopenfilename(title="选择原始 Excel", filetypes=[("Excel 文件", "*.xlsx *.xls")])
        if file_path:
            self.input_path_var.set(file_path)
            base_dir = os.path.dirname(file_path)
            if not self.dedup_output_dir_var.get().strip():
                self.dedup_output_dir_var.set(base_dir)
            if not self.batch_output_dir_var.get().strip():
                self.batch_output_dir_var.set(os.path.join(base_dir, "gemini_batches"))
            if not self.vertical_output_dir_var.get().strip():
                self.vertical_output_dir_var.set(base_dir)

    def choose_dedup_output_dir(self):
        d = filedialog.askdirectory(title="选择去重输出目录")
        if d:
            self.dedup_output_dir_var.set(d)

    def choose_batch_output_dir(self):
        d = filedialog.askdirectory(title="选择批量TXT输出目录")
        if d:
            self.batch_output_dir_var.set(d)

    def choose_vertical_output_dir(self):
        d = filedialog.askdirectory(title="选择纵向Excel输出目录")
        if d:
            self.vertical_output_dir_var.set(d)
            if not self.vertical_filename_var.get().strip():
                self.vertical_filename_var.set(self._default_vertical_filename())

    def choose_batch_source_txt(self):
        p = filedialog.askopenfilename(title="选择Batch源文件txt", filetypes=[("文本文件", "*.txt"), ("所有", "*.*")])
        if p:
            self.batch_source_txt_var.set(p)

    def _find_column(self, columns, target_name: str, case_insensitive: bool = False):
        if case_insensitive:
            return {str(c).strip().lower(): c for c in columns}.get(target_name.strip().lower())
        return {str(c).strip(): c for c in columns}.get(target_name.strip())

    def _load_required_columns(self, input_path: str) -> pd.DataFrame:
        df = pd.read_excel(input_path)
        if df.empty:
            raise ValueError("输入 Excel 为空，没有可处理的数据。")
        id_col = self._find_column(df.columns, "id", case_insensitive=True)
        title_col = self._find_column(df.columns, "英语-标题")
        desc_col = self._find_column(df.columns, "英语-描述")
        missing = [n for n, c in [("id", id_col), ("英语-标题", title_col), ("英语-描述", desc_col)] if c is None]
        if missing:
            raise ValueError(f"缺少必要列: {', '.join(missing)}")
        work_df = df[[id_col, title_col, desc_col]].copy()
        work_df.columns = ["id", "英语-标题", "英语-描述"]
        work_df["id"] = work_df["id"].astype(str).str.strip()
        work_df = work_df[work_df["id"] != ""].copy()
        work_df["英语-标题"] = work_df["英语-标题"].fillna("").astype(str)
        work_df["英语-描述"] = work_df["英语-描述"].fillna("").astype(str)
        return work_df

    def generate_gemini_batch_txt(self, dedup_df: pd.DataFrame, output_dir: str, batch_size: int):
        os.makedirs(output_dir, exist_ok=True)
        file_count = int(math.ceil(len(dedup_df) / batch_size))
        meta = []
        for i in range(file_count):
            chunk = dedup_df.iloc[i * batch_size: min((i + 1) * batch_size, len(dedup_df))]
            path = os.path.join(output_dir, f"gemini_batch_{i + 1:03d}.txt")
            ids = [str(v).strip() for v in chunk["id"].tolist()]
            blocks = []
            for _, row in chunk.iterrows():
                blocks.append("[PRODUCT]\n"
                              f"id: {str(row['id']).strip()}\n"
                              f"原始标题: {str(row['英语-标题']) if pd.notna(row['英语-标题']) else ''}\n"
                              f"原始描述: {str(row['英语-描述']) if pd.notna(row['英语-描述']) else ''}\n"
                              "[/PRODUCT]")
            with open(path, "w", encoding="utf-8") as f:
                f.write("IDS: " + ",".join(ids) + "\n\n" + "\n\n".join(blocks))
            meta.append({"file": path, "ids": ids})
        return file_count, meta

    def run_step1_pipeline(self):
        task = self._create_task("流程1：去重+切分", "原始Excel")
        try:
            input_path = self.input_path_var.get().strip()
            if not input_path or not os.path.exists(input_path):
                raise ValueError("请先选择有效的原始 Excel 文件。")
            batch_size = int(self.batch_size_var.get().strip() or "10")
            if batch_size <= 0:
                raise ValueError("批量大小必须是大于 0 的整数。")

            dedup_dir = self.dedup_output_dir_var.get().strip()
            batch_dir = self.batch_output_dir_var.get().strip()
            if not dedup_dir or not batch_dir:
                raise ValueError("请配置去重输出目录与批量TXT输出目录。")

            dedup_name = self._ensure_xlsx_name(self.dedup_filename_var.get())
            self.dedup_filename_var.set(dedup_name)
            dedup_path = os.path.join(dedup_dir, dedup_name)

            work_df = self._load_required_columns(input_path)
            dedup_df = work_df.drop_duplicates(subset=["id"], keep="first").reset_index(drop=True)
            if dedup_df.empty:
                raise ValueError("去重后没有可输出的数据。")

            os.makedirs(dedup_dir, exist_ok=True)
            dedup_df.to_excel(dedup_path, index=False)
            self._append_log(f"去重导出完成: {dedup_path}", dedup_path)

            file_count, batch_meta = self.generate_gemini_batch_txt(dedup_df, batch_dir, batch_size)
            self._append_log(f"批量TXT完成: {batch_dir}（共 {file_count} 个文件）", batch_dir)

            remark = json.dumps(batch_meta[:10], ensure_ascii=False)
            if len(batch_meta) > 10:
                remark += " ..."
            self._update_task(task, status="已完成", count=str(len(dedup_df)), output=dedup_path, remark=f"TXT共{file_count}个; {remark}")
            messagebox.showinfo("完成", f"流程1完成：\n去重文件：{dedup_path}\n批量目录：{batch_dir}")
        except Exception as exc:
            self._append_log(f"流程1失败: {exc}")
            self._update_task(task, status="失败", remark=str(exc))
            messagebox.showerror("错误", str(exc))

    def _parse_tsv_text(self, tsv_text: str) -> pd.DataFrame:
        cleaned = self.clean_gemini_code_block(tsv_text)
        if not cleaned.strip():
            raise ValueError("未检测到可解析的 TSV 内容。")
        cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
        if "\\n" in cleaned:
            cleaned = cleaned.replace("\\n", "\n")

        lines = [ln.rstrip() for ln in cleaned.split("\n") if ln.strip()]
        if not lines:
            raise ValueError("解析不到表头，请确认 TSV 文本不为空。")

        header_raw = lines[0].strip()
        headers = [self._normalize_col_name(h) for h in header_raw.split("\t")]
        self._append_log("读取到TSV表头: " + " | ".join(headers))

        if headers and headers[0] == "id品牌识别":
            self._append_log("检测到 Gemini 表头错误：id 和 品牌识别 缺少 Tab，请修正提示词后重新生成。正确表头应以 `id\t品牌识别\tAI_通用标题_EN...` 开头。")
            raise ValueError("检测到 Gemini 表头错误：id 和 品牌识别 缺少 Tab，请修正提示词后重新生成。")

        rows = []
        for line in lines[1:]:
            if line.strip() == header_raw:
                continue
            parts = line.split("\t")
            if len(parts) < len(headers):
                parts += [""] * (len(headers) - len(parts))
            elif len(parts) > len(headers):
                parts = parts[: len(headers) - 1] + ["\t".join(parts[len(headers) - 1 :])]
            rows.append(parts)

        if not rows:
            raise ValueError("没有解析到有效数据行，请检查 TSV 内容格式。")

        df = pd.DataFrame(rows, columns=headers).fillna("")
        for col in df.columns:
            df[col] = df[col].apply(lambda x: self.clean_export_text(self._clean_cell(x)))
        return df

    def read_gemini_result_file(self, input_path: str) -> pd.DataFrame:
        ext = os.path.splitext(input_path)[1].lower()
        if ext in {".txt", ".tsv"}:
            for enc in ("utf-8-sig", "utf-8", "gbk"):
                try:
                    with open(input_path, "r", encoding=enc) as f:
                        return self._parse_tsv_text(f.read())
                except Exception:
                    continue
            raise ValueError("读取 Gemini 结果文本失败，请检查文件编码。")
        if ext in {".xlsx", ".xls"}:
            df = pd.read_excel(input_path, dtype=str)
            if df.empty:
                raise ValueError("输入文件为空，没有可转换的数据。")
            df.columns = [self._normalize_col_name(c) for c in df.columns]
            self._append_log("读取到文件表头: " + " | ".join(df.columns))
            if df.columns[0] == "id品牌识别":
                raise ValueError("检测到 Gemini 表头错误：id 和 品牌识别 缺少 Tab，请修正提示词后重新生成。")
            df = df.fillna("")
            for col in df.columns:
                df[col] = df[col].apply(lambda x: self.clean_export_text(self._clean_cell(x)))
            return df
        raise ValueError("仅支持 .txt / .tsv / .xlsx / .xls 文件。")

    def parse_batch_source_txt(self, file_path) -> dict:
        text = None
        for enc in ("utf-8-sig", "utf-8", "gbk"):
            try:
                with open(file_path, "r", encoding=enc) as f:
                    text = f.read()
                break
            except Exception:
                continue
        if text is None:
            raise ValueError("Batch 源文件读取失败，请检查编码。")

        result = {}
        blocks = re.findall(r"\[PRODUCT\](.*?)\[/PRODUCT\]", text, re.S)
        for block in blocks:
            lines = block.strip("\n").split("\n")
            pid = ""
            title = ""
            desc = ""
            desc_started = False
            desc_lines = []
            for ln in lines:
                if ln.startswith("id:"):
                    pid = ln.split(":", 1)[1].strip()
                elif ln.startswith("原始标题:"):
                    title = ln.split(":", 1)[1].strip()
                elif ln.startswith("原始描述:"):
                    desc_started = True
                    desc_lines.append(ln.split(":", 1)[1].lstrip())
                elif desc_started:
                    desc_lines.append(ln)
            desc = "\n".join(desc_lines)
            if pid:
                result[pid] = {"原始标题": self.clean_export_text(title), "原始描述": self.clean_export_text(desc)}
        return result

    def _join_old_spec_fields(self, row: pd.Series, keys: list[str]) -> str:
        vals = [self._clean_cell(row.get(k, "")) for k in keys]
        return " || ".join([v for v in vals if v])

    def _format_spec_multiline(self, text: str) -> str:
        raw = self._clean_cell(text)
        if not raw:
            return ""

        # 先按 || 切段，再按换行展开，确保无论哪种来源都能清洗已有序号
        items = []
        for seg in re.split(r"\s*\|\|\s*", raw):
            for line in str(seg).splitlines():
                cleaned = self.clean_export_text(self.clean_leading_numbering(self.clean_export_text(line))).strip()
                if cleaned:
                    items.append(cleaned)

        return self.clean_export_text("\n".join(items))

    def _validate_and_normalize_new_fields(self, df: pd.DataFrame):
        missing = [f for f in REQUIRED_GEMINI_FIELDS if f not in df.columns]
        if missing:
            raise ValueError("缺少新版必需字段: " + ", ".join(missing))

    def _build_vertical_field_order(self, include_title: bool, include_desc: bool):
        base_fields = ["id"]
        if include_title:
            base_fields += ["原始标题"]
        if include_desc:
            base_fields += ["原始描述"]
        base_fields += BASE_VERTICAL_FIELDS[1:]
        return base_fields

    def convert_gemini_result_to_vertical_excel(self, df: pd.DataFrame, output_path: str, include_title: bool = False, include_desc: bool = False, source_map: dict | None = None):
        self._validate_and_normalize_new_fields(df)

        has_new_en = "AI_规格描述_EN" in df.columns
        has_new_cn = "AI_规格描述_CN" in df.columns
        has_old_en = any(c in df.columns for c in OLD_SPEC_EN_FIELDS)
        has_old_cn = any(c in df.columns for c in OLD_SPEC_CN_FIELDS)

        fields = self._build_vertical_field_order(include_title, include_desc)
        records = []
        matched = 0
        missing_original_ids = []

        for _, row in df.iterrows():
            rid = str(row.get("id", "")).strip()
            spec_en = self.clean_export_text(row.get("AI_规格描述_EN", "")) if has_new_en else self.clean_export_text(self._join_old_spec_fields(row, OLD_SPEC_EN_FIELDS))
            spec_cn = self.clean_export_text(row.get("AI_规格描述_CN", "")) if has_new_cn else self.clean_export_text(self._join_old_spec_fields(row, OLD_SPEC_CN_FIELDS))

            original_title = ""
            original_desc = ""
            if include_title or include_desc:
                if source_map and rid in source_map:
                    original_title = source_map[rid].get("原始标题", "")
                    original_desc = source_map[rid].get("原始描述", "")
                    matched += 1
                else:
                    original_title = row.get("原始标题", "") if "原始标题" in df.columns else ""
                    original_desc = row.get("原始描述", "") if "原始描述" in df.columns else ""
                    if not original_title and not original_desc:
                        missing_original_ids.append(rid)

            for field in fields:
                if field == "原始标题":
                    value = self.clean_export_text(original_title)
                elif field == "原始描述":
                    value = self.clean_export_text(original_desc)
                elif field == "AI_规格描述_EN":
                    value = self._format_spec_multiline(spec_en)
                elif field == "AI_规格描述_CN":
                    value = self._format_spec_multiline(spec_cn)
                else:
                    value = row.get(field, "")
                    if field in SELLING_POINT_FIELDS:
                        value = self.clean_export_text(value)
                    value = self.clean_leading_numbering(value)
                    value = self.clean_export_text(value)
                records.append({"字段": field, "内容": value})
            records.append({"字段": "", "内容": ""})

        out_df = pd.DataFrame(records, columns=["字段", "内容"])
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            out_df.to_excel(writer, index=False, sheet_name="Vertical")
            ws = writer.sheets["Vertical"]
            self.apply_vertical_excel_styles(ws)

        return matched, missing_original_ids

    def apply_vertical_excel_styles(self, ws):
        field_fill = PatternFill(fill_type="solid", fgColor="DCEBFF")
        id_fill = PatternFill(fill_type="solid", fgColor="C7DFFF")

        ws.freeze_panes = "A2"
        ws.cell(1, 1).font = Font(bold=True)
        ws.cell(1, 2).font = Font(bold=True)
        ws.cell(1, 1).fill = field_fill
        ws.cell(1, 2).fill = field_fill

        for i in range(2, ws.max_row + 1):
            fc = ws.cell(i, 1)
            cc = ws.cell(i, 2)
            fc.font = Font(bold=True)
            fc.fill = field_fill
            fc.alignment = Alignment(vertical="top")
            cc.alignment = Alignment(wrap_text=True, vertical="top")
            text_val = str(cc.value or "")
            logical_lines = text_val.split("\n") if text_val else [""]
            # 按列宽估算自动换行后的可视行数，避免卖点/描述显示不全
            visual_lines = 0
            for ln in logical_lines:
                visual_lines += max(1, math.ceil(len(ln) / 85))
            ws.row_dimensions[i].height = max(13.5, min(240, visual_lines * 13.5))
            if str(fc.value).strip() == "id":
                fc.fill = id_fill
                cc.fill = id_fill
                cc.font = Font(bold=True)

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 95

    def _build_vertical_output_path(self):
        out_dir = self.vertical_output_dir_var.get().strip()
        if not out_dir:
            raise ValueError("请先选择纵向Excel输出目录。")
        filename = self._ensure_xlsx_name(self.vertical_filename_var.get())
        self.vertical_filename_var.set(filename)
        return os.path.join(out_dir, filename)

    def _run_vertical(self, df: pd.DataFrame, source_type: str):
        task = self._create_task("Gemini TSV转纵向表", source_type)
        try:
            source_map = None
            remark_extra = ""
            include_title = self.include_original_title_var.get()
            include_desc = self.include_original_desc_var.get()
            if include_title or include_desc:
                txt = self.batch_source_txt_var.get().strip()
                if txt:
                    source_map = self.parse_batch_source_txt(txt)
                    remark_extra = f"; batch源文件: {txt}"
                else:
                    self._append_log("提示：未选择 batch 源文件 txt，将尝试从TSV旧字段回退原始标题/描述。")

            out_path = self._build_vertical_output_path()
            self._clean_stats = {"cite": 0, "tail_blank": 0, "wrap_quote": 0}
            matched, missing_ids = self.convert_gemini_result_to_vertical_excel(df, out_path, include_title, include_desc, source_map)

            if missing_ids:
                for mid in missing_ids[:20]:
                    self._append_log(f"未在源文件中找到 id: {mid} 的原始标题/描述")
            remark = f"转换成功{remark_extra}"
            include_title = self.include_original_title_var.get()
            include_desc = self.include_original_desc_var.get()
            if include_title or include_desc:
                remark += f"; 匹配{matched}个ID; 未匹配{len(missing_ids)}个ID"

            self._append_log(f"已清理 Gemini 引用标记：{self._clean_stats['cite']} 处")
            self._append_log(f"已清理尾部空行：{self._clean_stats['tail_blank']} 处")
            self._append_log(f"已清理包裹性双引号：{self._clean_stats['wrap_quote']} 处")
            self._append_log(f"转换完成，共处理 {len(df)} 个产品，输出文件：{out_path}", out_path)
            self._update_task(task, status="已完成", count=str(len(df)), output=out_path, remark=remark)
            messagebox.showinfo("完成", f"转换完成，共处理 {len(df)} 个产品。\n输出文件：\n{out_path}")
        except Exception as exc:
            self._append_log(f"Gemini 纵向转换失败: {exc}")
            self._update_task(task, status="失败", remark=str(exc))
            messagebox.showerror("错误", str(exc))

    def run_convert_gemini_vertical_from_paste(self):
        text = self.tsv_paste_text.get("1.0", "end")
        df = self._parse_tsv_text(text)
        self._run_vertical(df, "粘贴文本")


def main():
    if DND_AVAILABLE:
        root = TkinterDnD.Tk()
    else:
        root = tk.Tk()
    ExcelExtractorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
